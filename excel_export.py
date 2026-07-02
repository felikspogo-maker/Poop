"""Накопительная таблица участниц в формате Excel (.xlsx).

Каждая новая заявка дописывается отдельной строкой в один файл, так что
организатор получает единый список всех участниц.
"""

import logging
import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import config

logger = logging.getLogger(__name__)

SHEET_TITLE = "Участницы"

# Заголовки столбцов и соответствующие ключи из данных заявки
COLUMNS: list[tuple[str, str]] = [
    ("Дата заявки", "submitted_at"),
    ("ФИ", "name"),
    ("Город", "city"),
    ("Возраст", "age"),
    ("Рост, см", "height"),
    ("Семейное положение", "marital"),
    ("Категория", "category"),
    ("Телефон", "phone"),
    ("Опыт участия", "experience"),
    ("Фото, шт", "photos_count"),
    ("Telegram", "username"),
    ("Telegram ID", "telegram_id"),
    ("Согласие", "consent"),
]

# Ширина столбцов для читаемости
_WIDTHS = [18, 24, 16, 8, 9, 20, 34, 18, 30, 8, 18, 14, 22]


def _create_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    ws.append([title for title, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"  # закрепить строку заголовков
    return wb


def append_participant(data: dict) -> str | None:
    """Дописывает участницу в таблицу. Возвращает путь к файлу или None при ошибке."""
    path = config.EXCEL_FILE
    try:
        if os.path.exists(path):
            wb = load_workbook(path)
            ws = wb[SHEET_TITLE] if SHEET_TITLE in wb.sheetnames else wb.active
        else:
            wb = _create_workbook()
            ws = wb.active

        row = dict(data)
        row.setdefault("submitted_at", datetime.now().strftime("%d.%m.%Y %H:%M"))
        username = row.get("username")
        row["username"] = f"@{username}" if username else ""

        ws.append([row.get(key, "") for _, key in COLUMNS])
        wb.save(path)
        logger.info("Участница добавлена в таблицу: %s", path)
        return path
    except Exception:
        logger.exception("Не удалось записать участницу в таблицу Excel")
        return None
